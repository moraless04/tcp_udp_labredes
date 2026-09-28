#!/usr/bin/env python3
"""Gera monitor_dns_v2.docx preservando o documento institucional original."""

from __future__ import annotations

import hashlib
import shutil
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "monitor_dns.docx"
OUTPUT = ROOT / "monitor_dns_v2.docx"
WORKING = ROOT / ".docx-v2-work" / "monitor_dns_v2.generated.docx"
EXPECTED_SHA256 = "addcc001097701adba778a54ab10a48eb37423f22048d576d2a77d875dd32203"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def merge_numbering(original: bytes, generated: bytes) -> bytes:
    original_root = etree.fromstring(original)
    generated_root = etree.fromstring(generated)
    num_tag = qn("w:num")
    num_id_attr = qn("w:numId")
    original_ids = [int(node.get(num_id_attr)) for node in original_root.findall(num_tag)]
    original_max = max(original_ids, default=0)
    for node in generated_root.findall(num_tag):
        if int(node.get(num_id_attr)) > original_max:
            original_root.append(deepcopy(node))
    return etree.tostring(
        original_root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )


def compose_fidelity_package(source: Path, generated: Path, output: Path):
    temporary = output.with_suffix(".tmp.docx")
    with ZipFile(source) as src, ZipFile(generated) as gen, ZipFile(temporary, "w", ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "word/document.xml":
                data = gen.read(info.filename)
            elif info.filename == "word/numbering.xml":
                data = merge_numbering(data, gen.read(info.filename))
            elif info.filename == "docProps/core.xml":
                data = gen.read(info.filename)
            out.writestr(info, data)
    temporary.replace(output)


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:shd"))
    if old is not None:
        ppr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def set_bottom_border(paragraph, color="B7B7B7", size="4"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def add_num_id(doc: Document, fmt: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        lvl = abstract.find(qn("w:lvl"))
        if lvl is None:
            continue
        num_fmt = lvl.find(qn("w:numFmt"))
        if num_fmt is not None and num_fmt.get(qn("w:val")) == fmt:
            abstract_id = int(abstract.get(qn("w:abstractNumId")))
            break
    if abstract_id is None:
        raise RuntimeError(f"Nenhuma definição de numeração {fmt!r} encontrada")

    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int, level: int = 0):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def build():
    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("O documento original mudou; refaça a destilação antes de gerar a v2.")

    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    # O documento original é a autoridade visual. Preserva-se o bloco institucional,
    # a imagem ancorada, o título e a primeira faixa; o restante é reconstruído.
    paragraphs = list(doc.paragraphs)
    if paragraphs[5].text.strip() != "Monitoração do DNS":
        raise RuntimeError("Estrutura inesperada no título do documento original.")

    band_prototype = deepcopy(paragraphs[7]._p)
    keep = [p._p for p in paragraphs[:8]]
    body = doc._element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if not any(child is node for node in keep):
            body.remove(child)

    title = Paragraph(keep[5], doc._body)
    title.clear()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(2)
    set_font(title.add_run("Monitoração do DNS em Ambiente Podman"), size=12, bold=True)

    subtitle = Paragraph(keep[6], doc._body)
    subtitle.clear()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    set_font(subtitle.add_run("Versão 2 — Wireshark gráfico via noVNC"), size=10, italic=True)

    objective_band = Paragraph(keep[7], doc._body)
    objective_band.clear()
    set_font(objective_band.add_run("Objetivos"), name="Century Gothic", size=10, bold=True)

    def add_body(text=None, chunks=None, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        paragraph = doc.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = 1.08
        if chunks is None:
            chunks = [(text or "", False, False)]
        for value, bold, italic in chunks:
            set_font(paragraph.add_run(value), size=12, bold=bold, italic=italic)
        return paragraph

    def add_band(text):
        clone = deepcopy(band_prototype)
        body.insert(len(body) - 1, clone)
        paragraph = Paragraph(clone, doc._body)
        paragraph.clear()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(6)
        set_font(paragraph.add_run(text), name="Century Gothic", size=10, bold=True)
        return paragraph

    def add_heading(text, page_before=False):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.page_break_before = page_before
        set_font(paragraph.add_run(text), size=12, bold=True)
        return paragraph

    def add_command(command):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_together = True
        set_cell_shading(paragraph, "F2F2F2")
        set_bottom_border(paragraph, color="D9D9D9", size="2")
        set_font(paragraph.add_run(command), name="Courier New", size=9.5)
        return paragraph

    def add_note(label, text):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(7)
        set_cell_shading(paragraph, "E7E6E6")
        set_font(paragraph.add_run(label), size=11, bold=True)
        set_font(paragraph.add_run(text), size=11)
        return paragraph

    def add_list(items, ordered=False):
        num_id = add_num_id(doc, "decimal" if ordered else "bullet")
        for item in items:
            paragraph = doc.add_paragraph()
            apply_num(paragraph, num_id)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.05
            if isinstance(item, str):
                chunks = [(item, False, False)]
            else:
                chunks = item
            for value, bold, italic in chunks:
                set_font(paragraph.add_run(value), size=12, bold=bold, italic=italic)

    add_body("Ao concluir este exercício, você deverá ser capaz de:")
    add_list([
        "Iniciar, verificar, parar e limpar um laboratório DNS construído com contêineres Podman.",
        "Gerar consultas DNS controladas com nslookup e dig, incluindo diferentes tipos de registro.",
        "Capturar e localizar o tráfego DNS no Wireshark por meio de uma interface gráfica noVNC.",
        "Identificar a estrutura de uma mensagem DNS: cabeçalho, pergunta, resposta e seções adicionais.",
        "Reconhecer registros A, AAAA, MX, NS, CNAME, SOA, PTR e TXT.",
        "Interpretar Transaction ID, flags, TTL, EDNS0 e o protocolo de transporte utilizado.",
    ])

    add_band("Descrição")
    add_heading("Arquitetura do ambiente")
    add_body(
        "O laboratório utiliza uma rede bridge privada chamada dns-lab-net (10.53.0.0/24) e três contêineres. "
        "O cliente e o monitor pertencem ao mesmo pod, dns-station, e compartilham a interface eth0. "
        "Essa decisão permite que o Wireshark observe exatamente os pacotes emitidos pelo nslookup."
    )
    add_list([
        [("dns-client: ", True, False), ("terminal do aluno com nslookup, dig e ferramentas de rede.", False, False)],
        [("dns-monitor: ", True, False), ("Wireshark, desktop Openbox, Xvfb, x11vnc e noVNC.", False, False)],
        [("dns-resolver: ", True, False), ("Unbound em 10.53.0.53, usado como resolvedor padrão com cache controlável.", False, False)],
        [("Host: ", True, False), ("publica somente 127.0.0.1:6080 para o navegador; a porta DNS não é exposta.", False, False)],
    ])
    add_note(
        "Importante: ",
        "os resultados DNS, endereços e valores de TTL mudam com o tempo. Registre o que for observado durante a sua execução.",
    )

    add_heading("Requisitos no host")
    add_list([
        "Podman 5 ou superior e uma Podman Machine Linux em execução no macOS ou Windows; em Linux, apenas o serviço Podman é necessário.",
        "Navegador com acesso a http://127.0.0.1:6080/vnc.html.",
        "Acesso externo a UDP e TCP porta 53. VPNs ou redes institucionais podem bloquear consultas diretas a resolvedores públicos.",
        "Terminal aberto no diretório que contém este documento e o script scripts/lab.sh.",
    ])

    add_band("Preparação e acesso ao ambiente")
    add_heading("1. Iniciar a Podman Machine")
    add_body("No macOS ou Windows, confirme que a máquina virtual do Podman está em execução:")
    add_command("podman machine start")
    add_body("Se a máquina já estiver iniciada, o Podman apenas informará esse estado.")

    add_heading("2. Construir e iniciar o laboratório")
    add_body("No diretório do laboratório, execute:")
    add_command("./scripts/lab.sh up")
    add_note(
        "Primeira execução: ",
        "as três imagens serão construídas. A imagem gráfica contém Wireshark e bibliotecas Qt, portanto o primeiro build pode levar alguns minutos.",
    )
    add_body("Ao final, o script apresenta a URL do noVNC e o comando para abrir o terminal do aluno.")

    add_heading("3. Verificar o estado e a conectividade")
    add_command("./scripts/lab.sh status")
    add_command("./scripts/lab.sh doctor")
    add_body(
        "O diagnóstico esperado informa Podman OK, noVNC OK e acesso UDP/53 aos resolvedores 8.8.8.8 e 1.1.1.1. "
        "Se um desses testes falhar, consulte a seção Solução de problemas antes de iniciar a Parte D."
    )

    add_heading("4. Acessar o Wireshark pelo navegador")
    add_list([
        "Abra http://127.0.0.1:6080/vnc.html.",
        "Clique em Connect.",
        "Confirme que o Wireshark está aberto e capturando na interface eth0 com o filtro de captura udp port 53 or tcp port 53.",
        "Use o filtro de exibição dns para mostrar apenas mensagens DNS.",
    ], ordered=True)
    add_note(
        "Acesso local sem senha: ",
        "a porta 6080 é publicada exclusivamente em 127.0.0.1. Não altere o mapeamento para 0.0.0.0 sem adicionar autenticação e controles de acesso.",
    )

    add_heading("5. Abrir o terminal do cliente")
    add_body("Mantenha o navegador aberto. Em outro terminal do host, execute:")
    add_command("./scripts/lab.sh shell")
    add_body("Dentro do contêiner, confirme o resolvedor padrão:")
    add_command("cat /etc/resolv.conf")
    add_body("O resultado esperado contém nameserver 10.53.0.53. Para sair do cliente, use exit.")

    add_heading("6. Limpar o cache entre consultas")
    add_body(
        "O cliente não utiliza systemd-resolved. Antes de consultas que devem partir de um cache vazio, execute no terminal do host:"
    )
    add_command("./scripts/lab.sh flush-dns")
    add_body(
        "Esse comando limpa o cache do Unbound. Não é possível limpar o cache dos resolvedores públicos 8.8.8.8 ou 1.1.1.1."
    )

    add_band("Execução dos experimentos")
    add_note(
        "Organização recomendada: ",
        "use um terminal do host para comandos ./scripts/lab.sh e outro terminal conectado ao dns-client. O Wireshark permanece aberto no navegador.",
    )

    add_heading("Parte A — Consulta básica (registro A)")
    add_list([
        "No Wireshark, limpe a lista de pacotes ou reinicie a captura em eth0.",
        "No host, limpe o cache do resolvedor do laboratório.",
        "No dns-client, execute a consulta A abaixo.",
        "Pare a captura e localize o par consulta/resposta usando o filtro dns.",
    ], ordered=True)
    add_command("./scripts/lab.sh flush-dns        # executar no host")
    add_command("nslookup -type=A www.google.com  # executar no dns-client")
    add_heading("Responda:")
    add_list([
        "Qual endereço IPv4 foi retornado? Ele aparece tanto no nslookup quanto na captura?",
        "Qual protocolo de transporte foi utilizado e qual é a porta do servidor DNS?",
        "Compare o Transaction ID da consulta e da resposta. Para que ele serve?",
        "Na seção Queries, identifique nome, Type e Class.",
        "Na seção Answers, identifique tipo, endereço e TTL. Explique o significado prático do TTL.",
        "Confirme que os pacotes foram enviados entre 10.53.0.10 e 10.53.0.53.",
    ], ordered=True)

    add_heading("Parte B — Diferentes tipos de registro")
    add_body(
        "Para cada consulta abaixo, limpe o cache no host, reinicie a captura e execute o comando no dns-client. "
        "Registre a saída e os campos correspondentes no Wireshark."
    )
    add_command("nslookup -type=A     www.wikipedia.org")
    add_command("nslookup -type=AAAA  www.google.com")
    add_command("nslookup -type=MX    gmail.com")
    add_command("nslookup -type=NS    ufrgs.br")
    add_command("nslookup -type=SOA   ufrgs.br")
    add_command("nslookup -type=TXT   google.com")
    add_command("nslookup -type=CNAME www.wikipedia.org")
    add_heading("Responda:")
    add_list([
        "Explique em uma frase o que representa cada tipo A, AAAA, MX, NS, SOA, TXT e CNAME, e registre o que foi retornado.",
        "Qual é a diferença entre A e AAAA? A que versão do IP cada tipo se refere?",
        "Nos registros MX, o que significa o valor de preferência e como ele ordena os servidores?",
        "O que o NS revela sobre ufrgs.br? No SOA, identifique o servidor primário e ao menos um temporizador.",
        "Que informações apareceram em TXT? Há políticas de e-mail ou registros de verificação?",
        "O CNAME retornou um nome canônico? Verifique se houve outra consulta para resolver esse nome.",
    ], ordered=True)

    add_heading("Parte C — Consulta reversa (registro PTR)")
    add_body("Limpe o cache, reinicie a captura e execute no dns-client:")
    add_command("nslookup -type=PTR 8.8.8.8")
    add_heading("Responda:")
    add_list([
        "Que nome foi retornado para 8.8.8.8?",
        "Como o endereço aparece na seção Queries? Explique os octetos invertidos e o sufixo in-addr.arpa.",
        "Uma consulta reversa bem-sucedida garante que a consulta direta retorne o mesmo endereço? Discuta brevemente.",
    ], ordered=True)

    add_heading("Parte D — Servidor específico e modo interativo")
    add_body("Reinicie a captura. As consultas abaixo ignoram o Unbound e vão diretamente aos resolvedores informados:")
    add_command("nslookup www.wikipedia.org 8.8.8.8")
    add_command("nslookup www.wikipedia.org 1.1.1.1")
    add_body("Depois, entre no modo interativo do nslookup e execute:")
    add_command("nslookup")
    add_command("> server 8.8.8.8")
    add_command("> set type=MX")
    add_command("> gmail.com")
    add_command("> set type=NS")
    add_command("> wikipedia.org")
    add_command("> exit")
    add_heading("Responda:")
    add_list([
        "Confirme no Wireshark os destinos 8.8.8.8 e 1.1.1.1.",
        "As respostas foram equivalentes? Compare também os TTLs e relacione diferenças ao estado do cache de cada servidor.",
        "No modo interativo, o comando server alterou o destino das consultas seguintes? Comprove pela captura.",
    ], ordered=True)

    add_heading("Parte E — Cabeçalho, flags, transporte e EDNS0")
    add_body("Selecione uma resposta DNS no Wireshark, expanda Domain Name System e examine Flags.")
    add_heading("Responda:")
    add_list([
        "RD: quem define esse bit e o que ele solicita?",
        "RA: quem define esse bit e o que ele informa?",
        "AA: o que indica quando está ligado?",
        "O RD estava ligado? Relacione consulta recursiva e resolução iterativa.",
    ], ordered=True)
    add_body("Para produzir deliberadamente uma consulta DNS sobre TCP, execute no dns-client:")
    add_command("dig +tcp @10.53.0.53 www.google.com A")
    add_body("Use o filtro tcp.port == 53 e compare a troca TCP com a consulta UDP.")
    add_body("Para observar uma resposta maior e o registro OPT/EDNS0, execute:")
    add_command("dig +dnssec +bufsize=1232 @10.53.0.53 . DNSKEY")
    add_heading("Responda:")
    add_list([
        "Qual foi o tamanho da resposta? Ela ultrapassou o limite clássico de 512 bytes?",
        "Existe um registro OPT na seção Additional? Que tamanho UDP foi anunciado?",
        "Em que situações o DNS usa TCP, além do teste forçado? Considere truncamento e transferências de zona.",
    ], ordered=True)

    add_band("Salvar resultados")
    add_body(
        "No Wireshark, use File > Save As e salve a captura em /captures com extensão .pcapng, por exemplo /captures/dns-parte-a.pcapng. "
        "O diretório está montado na pasta captures/ do laboratório e permanece disponível no host."
    )
    add_note(
        "Antes de substituir um arquivo: ",
        "confira o nome da parte e preserve capturas que serão usadas no relatório. Os comandos de limpeza do laboratório não removem captures/.",
    )

    add_band("Parar e limpar o ambiente")
    add_heading("Parada temporária")
    add_body("Para interromper o laboratório e preservar pod e contêineres:")
    add_command("./scripts/lab.sh stop")
    add_body("Para reiniciar posteriormente:")
    add_command("./scripts/lab.sh start")

    add_heading("Remover somente pod e contêineres")
    add_command("./scripts/lab.sh down")
    add_body("As imagens, a rede dns-lab-net e captures/ são preservadas. Um novo up recria rapidamente os contêineres.")

    add_heading("Limpeza completa da infraestrutura local")
    add_command("./scripts/lab.sh clean")
    add_body(
        "Esse comando remove pod, contêineres, rede e as três imagens locais do laboratório. O diretório captures/ continua preservado e só deve ser apagado após revisão explícita dos arquivos."
    )

    add_band("Solução de problemas")
    add_list([
        [("Podman indisponível: ", True, False), ("execute podman machine start e depois ./scripts/lab.sh doctor.", False, False)],
        [("noVNC não abre: ", True, False), ("verifique ./scripts/lab.sh status e podman logs dns-monitor.", False, False)],
        [("Wireshark sem pacotes: ", True, False), ("confirme captura em eth0, filtro udp port 53 or tcp port 53 e nameserver 10.53.0.53.", False, False)],
        [("Resolvedor sem resposta: ", True, False), ("consulte podman logs dns-resolver e repita ./scripts/lab.sh flush-dns.", False, False)],
        [("8.8.8.8 ou 1.1.1.1 sem resposta: ", True, False), ("a rede ou VPN pode bloquear DNS externo; a Parte D depende da liberação de UDP/TCP porta 53.", False, False)],
        [("Conflito na porta 6080: ", True, False), ("encerre o serviço que ocupa 127.0.0.1:6080 antes de executar up.", False, False)],
    ])

    add_band("Referências")
    add_list([
        [("RFC 1034 e RFC 1035 — ", False, False), ("Domain Names: Concepts and Facilities / Implementation and Specification", False, True)],
        [("RFC 6891 — ", False, False), ("Extension Mechanisms for DNS (EDNS0)", False, True)],
        "Documentação do Podman: pods, redes bridge e execução rootless.",
        "Wireshark User's Guide e Display Filter Reference.",
        "Unbound documentation e unbound-control.",
    ])

    # Mantém a geometria A4 e as margens do original de forma explícita.
    section = doc.sections[0]
    section.left_margin = Inches(1.18)
    section.right_margin = Inches(1.18)
    section.top_margin = Inches(0.98)
    section.bottom_margin = Inches(0.98)

    for drawing_property in doc._element.iter(qn("wp:docPr")):
        drawing_property.set("title", "Brasão da PUCRS")
        drawing_property.set("descr", "Brasão da Pontifícia Universidade Católica do Rio Grande do Sul")

    doc.core_properties.title = "Monitoração do DNS em Ambiente Podman — versão 2"
    doc.core_properties.subject = "Laboratório de DNS com Wireshark via noVNC"
    doc.core_properties.comments = "Versão adaptada ao ambiente conteinerizado Podman."
    WORKING.parent.mkdir(parents=True, exist_ok=True)
    doc.save(WORKING)
    compose_fidelity_package(SOURCE, WORKING, OUTPUT)

    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("O documento original foi alterado durante a geração.")

    print(OUTPUT)


if __name__ == "__main__":
    build()
