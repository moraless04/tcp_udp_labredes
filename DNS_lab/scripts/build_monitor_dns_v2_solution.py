#!/usr/bin/env python3
"""Gera o relatório de solução dos experimentos de monitor_dns_v2.docx."""

from __future__ import annotations

import hashlib
import shutil
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "monitor_dns_v2.docx"
OUTPUT = ROOT / "monitor_dns_v2_solution.docx"
WORKING = ROOT / ".solution-work" / "monitor_dns_v2_solution.generated.docx"
EXPECTED_SHA256 = "9b42e9e36ab4f52d8775a801301cdf06b77599d516fcbf03cb06e0fccd5f69b1"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def merge_numbering(original: bytes, generated: bytes) -> bytes:
    original_root = etree.fromstring(original)
    generated_root = etree.fromstring(generated)
    num_tag = qn("w:num")
    num_id_attr = qn("w:numId")
    original_ids = [int(node.get(num_id_attr)) for node in original_root.findall(num_tag)]
    original_max = max(original_ids, default=0)
    for node in generated_root.findall(num_tag):
        if int(node.get(num_id_attr)) > original_max:
            original_root.append(deepcopy(node))
    return etree.tostring(
        original_root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )


def compose_fidelity_package(source: Path, generated: Path, output: Path):
    temporary = output.with_suffix(".tmp.docx")
    with ZipFile(source) as src, ZipFile(generated) as gen, ZipFile(
        temporary, "w", ZIP_DEFLATED
    ) as out:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "word/document.xml":
                data = gen.read(info.filename)
            elif info.filename == "word/numbering.xml":
                data = merge_numbering(data, gen.read(info.filename))
            elif info.filename == "docProps/core.xml":
                data = gen.read(info.filename)
            out.writestr(info, data)
    temporary.replace(output)


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_shading(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:shd"))
    if old is not None:
        ppr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def set_bottom_border(paragraph, color="D9D9D9", size="2"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def add_num_id(doc: Document, fmt: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        level = abstract.find(qn("w:lvl"))
        if level is None:
            continue
        num_fmt = level.find(qn("w:numFmt"))
        if num_fmt is not None and num_fmt.get(qn("w:val")) == fmt:
            abstract_id = int(abstract.get(qn("w:abstractNumId")))
            break
    if abstract_id is None:
        raise RuntimeError(f"Definição de numeração ausente: {fmt}")
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    current = ppr.find(qn("w:numPr"))
    if current is not None:
        ppr.remove(current)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def build():
    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("O documento de referência mudou; refaça a destilação.")

    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    paragraphs = list(doc.paragraphs)
    if paragraphs[5].text.strip() != "Monitoração do DNS em Ambiente Podman":
        raise RuntimeError("Estrutura inesperada no documento de referência.")

    band_prototype = deepcopy(paragraphs[7]._p)
    keep = [paragraph._p for paragraph in paragraphs[:8]]
    body = doc._element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if not any(child is node for node in keep):
            body.remove(child)

    title = Paragraph(keep[5], doc._body)
    title.clear()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(2)
    set_font(title.add_run("Relatório de solução — Monitoração do DNS"), size=12, bold=True)

    subtitle = Paragraph(keep[6], doc._body)
    subtitle.clear()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    set_font(
        subtitle.add_run("Resultados observados em 24 de agosto de 2026"),
        size=10,
        italic=True,
    )

    first_band = Paragraph(keep[7], doc._body)
    first_band.clear()
    set_font(first_band.add_run("Resumo da execução"), name="Century Gothic", size=10, bold=True)

    def add_body(text=None, chunks=None, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        paragraph = doc.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = 1.08
        if chunks is None:
            chunks = [(text or "", False, False)]
        for value, bold, italic in chunks:
            set_font(paragraph.add_run(value), size=12, bold=bold, italic=italic)
        return paragraph

    def add_band(text, page_before=False):
        clone = deepcopy(band_prototype)
        body.insert(len(body) - 1, clone)
        paragraph = Paragraph(clone, doc._body)
        paragraph.clear()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.page_break_before = page_before
        set_font(paragraph.add_run(text), name="Century Gothic", size=10, bold=True)
        return paragraph

    def add_heading(text, page_before=False):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(9)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.page_break_before = page_before
        set_font(paragraph.add_run(text), size=12, bold=True)
        return paragraph

    def add_command(command):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_together = True
        set_shading(paragraph, "F2F2F2")
        set_bottom_border(paragraph)
        for index, line in enumerate(command.split("\n")):
            if index:
                paragraph.add_run().add_break()
            set_font(paragraph.add_run(line), name="Courier New", size=9.2)
        return paragraph

    def add_note(label, text):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(7)
        set_shading(paragraph, "E7E6E6")
        set_font(paragraph.add_run(label), size=11, bold=True)
        set_font(paragraph.add_run(text), size=11)
        return paragraph

    def add_list(items, ordered=False):
        num_id = add_num_id(doc, "decimal" if ordered else "bullet")
        for item in items:
            paragraph = doc.add_paragraph()
            apply_num(paragraph, num_id)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.05
            chunks = [(item, False, False)] if isinstance(item, str) else item
            for value, bold, italic in chunks:
                set_font(paragraph.add_run(value), size=12, bold=bold, italic=italic)

    def add_answer(label, text):
        add_body(
            chunks=[(label, True, False), (text, False, False)],
            after=5,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )

    add_body(
        "Este relatório apresenta a solução das Partes A a E do experimento de monitoração DNS. "
        "Os comandos foram executados no laboratório Podman e as mensagens foram confirmadas por capturas na interface eth0 do pod dns-station."
    )
    add_note(
        "Escopo temporal: ",
        "a coleta ocorreu em 24/08/2026, por volta de 16:45 BRT (19:45 UTC). Endereços, TTLs, servidores e conteúdos TXT podem mudar em outras execuções.",
    )

    add_heading("Ambiente utilizado")
    add_list([
        [("Podman: ", True, False), ("cliente e servidor 6.0.2.", False, False)],
        [("Estação: ", True, False), ("dns-client e dns-monitor no endereço 10.53.0.10, compartilhando eth0.", False, False)],
        [("Resolvedor: ", True, False), ("Unbound em 10.53.0.53; /etc/resolv.conf contém nameserver 10.53.0.53.", False, False)],
        [("Monitoramento: ", True, False), ("Wireshark via noVNC em 127.0.0.1:6080; endpoint respondeu HTTP 200.", False, False)],
    ])
    add_body(
        "Neste laboratório, o Unbound ocupa a função de resolvedor e cache DNS do cliente, sem systemd-resolved. "
        "Ele está configurado como encaminhador com cache para resolvedores públicos, o que torna a limpeza do cache reproduzível com ./scripts/lab.sh flush-dns."
    )

    add_heading("Síntese dos resultados")
    add_list([
        "As consultas DNS padrão usaram UDP/53 e mantiveram o mesmo Transaction ID na consulta e na resposta.",
        "Foram identificados registros A, AAAA, MX, NS, SOA, TXT, CNAME e PTR.",
        "Consultas diretas a 8.8.8.8 e 1.1.1.1 retornaram o mesmo IPv4, mas TTLs e parâmetros EDNS diferentes.",
        "O teste +tcp produziu uma sessão TCP completa; a consulta DNSKEY produziu 1139 bytes sobre UDP com EDNS0.",
        "As capturas das Partes A, D e E registraram zero pacotes descartados.",
    ])

    add_band("Parte A — Consulta básica (registro A)", page_before=True)
    add_command("./scripts/lab.sh flush-dns\nnslookup -type=A www.google.com")
    add_heading("Resultado observado")
    add_command(
        "www.google.com.  297  IN  A  142.251.150.119 ... 142.251.157.119\n"
        "8 respostas A; consulta entre 10.53.0.10 e 10.53.0.53"
    )
    add_answer(
        "Endereços retornados: ",
        "142.251.150.119, 142.251.151.119, 142.251.152.119, 142.251.153.119, 142.251.154.119, 142.251.155.119, 142.251.156.119 e 142.251.157.119. Os mesmos valores apareceram no nslookup e na resposta capturada.",
    )
    add_answer(
        "Transporte e portas: ",
        "UDP. O pacote saiu de 10.53.0.10:55873 para 10.53.0.53:53; a resposta percorreu o caminho inverso. A porta 55873 é efêmera e 53 é a porta DNS do servidor.",
    )
    add_answer(
        "Transaction ID: ",
        "3993, equivalente a 0x0f99, nos dois pacotes. O identificador permite associar cada resposta à consulta que a originou, inclusive quando há várias consultas simultâneas.",
    )
    add_answer(
        "Seção Queries: ",
        "nome www.google.com, tipo A, classe IN. A classe IN representa a Internet e o tipo A solicita endereços IPv4.",
    )
    add_answer(
        "Seção Answers e TTL: ",
        "oito registros A com TTL de 297 segundos. O TTL indica por quanto tempo um cache pode reutilizar a resposta. Em duas leituras posteriores, o Unbound mostrou 146 e 144 segundos, confirmando a contagem regressiva do cache.",
    )
    add_answer(
        "Flags: ",
        "a consulta apresentou RD=1, solicitando recursão. A resposta apresentou QR=1, RD=1 e RA=1. AA ficou desligado porque o Unbound respondeu como resolvedor recursivo/encaminhador, não como servidor autoritativo de google.com.",
    )
    add_note(
        "Latência da amostra: ",
        "aproximadamente 25,8 ms entre consulta e resposta. Esse valor mede esta execução e não deve ser tratado como característica permanente do DNS.",
    )

    add_band("Parte B — Diferentes tipos de registro")
    add_heading("A e CNAME — www.wikipedia.org")
    add_command(
        "www.wikipedia.org.    86344  IN  CNAME  dyna.wikimedia.org.\n"
        "dyna.wikimedia.org.     180  IN  A      195.200.68.224"
    )
    add_body(
        "O tipo A associa um nome a um endereço IPv4. CNAME cria um alias para um nome canônico. "
        "A resposta mostrou que www.wikipedia.org é um alias de dyna.wikimedia.org e incluiu o registro A necessário para alcançar 195.200.68.224. "
        "Na consulta específica por CNAME, o TTL observado foi 86400 segundos."
    )

    add_heading("AAAA — www.google.com")
    add_command(
        "www.google.com.  300  IN  AAAA  2001:4860:4826:7700::\n"
        "... oito respostas no prefixo 2001:4860"
    )
    add_body(
        "AAAA associa um nome a IPv6, enquanto A associa a IPv4. Foram retornados oito endereços IPv6, todos com TTL 300. "
        "A coexistência de A e AAAA permite que clientes escolham a versão de IP compatível com sua conectividade."
    )

    add_heading("MX — gmail.com")
    add_command(
        "5  gmail-smtp-in.l.google.com.\n10 alt1.gmail-smtp-in.l.google.com.\n"
        "20 alt2...  30 alt3...  40 alt4..."
    )
    add_body(
        "MX identifica os servidores que recebem e-mail para o domínio. O menor número tem maior preferência; portanto, o servidor de preferência 5 é tentado antes dos servidores 10, 20, 30 e 40. "
        "O TTL observado foi 2833 segundos. Servidores de maior valor funcionam como alternativas quando os preferidos estão indisponíveis."
    )

    add_heading("NS e SOA — ufrgs.br")
    add_command(
        "NS: ns1.ufrgs.br.; ns2.ufrgs.br.; pampa.tche.br.\n"
        "SOA: ns1.ufrgs.br. dnsadmin.ufrgs.br. 2026082100 1200 3600 1209600 600"
    )
    add_body(
        "NS identifica os servidores autoritativos da zona. O SOA descreve a autoridade inicial: servidor primário ns1.ufrgs.br, contato dnsadmin@ufrgs.br, serial 2026082100, refresh 1200 s, retry 3600 s, expire 1209600 s e mínimo/negative TTL 600 s. "
        "Esses temporizadores orientam a sincronização de servidores secundários e o cache de respostas negativas."
    )

    add_heading("TXT — google.com")
    add_body(
        "Foram retornados 16 registros TXT, TTL 255. Entre eles estavam a política SPF v=spf1 include:_spf.google.com ~all e tokens de verificação de Google, Apple, Facebook, Microsoft, Cisco, DocuSign e OneTrust. "
        "TXT transporta texto associado ao domínio; na prática, é muito usado para políticas de e-mail, comprovação de domínio e integração de serviços."
    )
    add_note(
        "Interpretação: ",
        "um registro TXT não executa uma política por si só. O significado depende do consumidor, como um servidor de e-mail ao interpretar SPF ou um provedor ao validar um token.",
    )

    add_band("Parte C — Consulta reversa (registro PTR)")
    add_command("nslookup -type=PTR 8.8.8.8")
    add_command("8.8.8.8.in-addr.arpa.  81346  IN  PTR  dns.google.")
    add_answer(
        "Nome retornado: ",
        "dns.google. O registro PTR relaciona um endereço IPv4 a um nome, fazendo o caminho inverso de um registro A.",
    )
    add_answer(
        "Nome da consulta: ",
        "8.8.8.8.in-addr.arpa. Para IPv4, os octetos aparecem em ordem inversa sob a árvore especial in-addr.arpa; neste caso, a repetição dos quatro octetos torna a inversão visualmente indistinguível.",
    )
    add_answer(
        "Relação com consulta direta: ",
        "um PTR bem-sucedido não garante que uma consulta A ao nome retorne o mesmo endereço. DNS direto e reverso são zonas administradas separadamente e podem representar relações um-para-muitos ou configurações divergentes.",
    )

    add_band("Parte D — Servidor específico e modo interativo")
    add_command(
        "nslookup www.wikipedia.org 8.8.8.8\n"
        "nslookup www.wikipedia.org 1.1.1.1"
    )
    add_heading("Comparação dos resolvedores")
    add_answer(
        "Google Public DNS (8.8.8.8): ",
        "CNAME dyna.wikimedia.org com TTL 16555 e A 195.200.68.224 com TTL 106. A resposta anunciou tamanho EDNS UDP 512.",
    )
    add_answer(
        "Cloudflare (1.1.1.1): ",
        "o mesmo CNAME com TTL 86322 e o mesmo A 195.200.68.224 com TTL 102. A resposta anunciou tamanho EDNS UDP 1232.",
    )
    add_body(
        "As respostas foram semanticamente equivalentes porque produziram o mesmo alias e endereço final. Os TTLs diferiram porque cada resolvedor possuía um instante próprio de preenchimento do cache e possivelmente recebeu dados com tempos residuais diferentes. "
        "A captura confirmou destinos diretos 8.8.8.8:53 e 1.1.1.1:53, sem passar pelo Unbound."
    )

    add_heading("Modo interativo")
    add_command(
        "> server 8.8.8.8\n> set type=MX\n> gmail.com\n"
        "> set type=NS\n> wikipedia.org\n> exit"
    )
    add_body(
        "Após server 8.8.8.8, o nslookup informou 8.8.8.8 como servidor padrão. A consulta MX retornou os cinco servidores do Gmail com preferências 5, 10, 20, 30 e 40. "
        "A consulta NS retornou ns0.wikimedia.org, ns1.wikimedia.org e ns2.wikimedia.org. Portanto, a mudança de servidor permaneceu ativa para as consultas seguintes."
    )
    add_note(
        "Variabilidade externa: ",
        "uma tentativa automatizada de NS para 8.8.8.8 expirou duas vezes antes de obter resposta. A repetição interativa respondeu normalmente. Isso mostra que perda, filtragem ou atraso momentâneo podem afetar uma execução sem invalidar o procedimento.",
    )

    add_band("Parte E — Cabeçalho, flags, transporte e EDNS0")
    add_heading("Interpretação das flags")
    add_list([
        [("RD — Recursion Desired: ", True, False), ("definida pelo cliente; solicita que o servidor entregue uma resposta final, realizando ou encaminhando a resolução necessária.", False, False)],
        [("RA — Recursion Available: ", True, False), ("definida pelo servidor; informa que ele oferece o serviço recursivo. O Unbound respondeu com RA=1.", False, False)],
        [("AA — Authoritative Answer: ", True, False), ("indica que a resposta veio de um servidor autoritativo para o nome consultado. Ficou desligada nas respostas recursivas observadas.", False, False)],
        [("QR — Query/Response: ", True, False), ("distingue consulta de resposta; QR=1 apareceu nas respostas.", False, False)],
    ])
    add_body(
        "O RD estava ligado. Do ponto de vista do cliente, a consulta foi recursiva: ele pediu uma resposta completa ao Unbound. "
        "O Unbound, configurado como encaminhador com cache, consultou resolvedores públicos; estes executam a resolução necessária, que internamente pode envolver consultas iterativas aos servidores raiz, TLD e autoritativos."
    )

    add_heading("DNS sobre TCP")
    add_command("dig +tcp @10.53.0.53 www.google.com A")
    add_body(
        "A captura mostrou SYN, SYN-ACK e ACK, seguidos da consulta e da resposta na conexão entre 10.53.0.10:40947 e 10.53.0.53:53. "
        "O Transaction ID foi 51655, a resposta DNS teve 171 bytes, retornou oito registros A com TTL 170 e apresentou flags qr rd ra."
    )
    add_body(
        "Além do uso forçado por +tcp, DNS utiliza TCP quando uma resposta UDP vem truncada e precisa ser repetida, em transferências de zona AXFR/IXFR e em outros fluxos que exigem transporte confiável ou excedem o tamanho UDP negociado."
    )

    add_heading("EDNS0 e resposta DNSKEY")
    add_command("dig +dnssec +bufsize=1232 @10.53.0.53 . DNSKEY")
    add_answer(
        "Tamanho: ",
        "1139 bytes de mensagem DNS e 1167 bytes no pacote IP. O resultado ultrapassou o limite DNS/UDP clássico de 512 bytes.",
    )
    add_answer(
        "OPT/EDNS0: ",
        "presente na seção Additional, versão 0, com tamanho UDP anunciado de 1232 bytes e bit DO ligado. O DO solicita dados DNSSEC.",
    )
    add_answer(
        "Transporte: ",
        "UDP foi suficiente porque 1139 bytes cabem no limite EDNS0 de 1232. A resposta continha três DNSKEY e um RRSIG, totalizando quatro registros na seção Answer.",
    )

    add_band("Evidências, limitações e conclusão")
    add_heading("Resumo das capturas")
    add_list([
        "Parte A: dois pacotes — uma consulta e uma resposta UDP — sem descartes.",
        "Parte D: oito pacotes capturados, confirmando destinos 8.8.8.8 e 1.1.1.1, sem descartes.",
        "Parte E: doze pacotes, incluindo a sessão TCP e o par UDP DNSKEY/EDNS0, sem descartes.",
    ])
    add_note(
        "Checksums na captura: ",
        "o tcpdump no host marcou alguns checksums como incorretos. Isso decorre do checksum offload na pilha virtual: a captura ocorre antes de o checksum ser finalizado. Como houve respostas válidas e zero descarte no dumpcap, não se tratou de corrupção dos pacotes.",
    )

    add_heading("Limitações")
    add_list([
        "DNS é dinâmico: endereços, TTLs, registros TXT e distribuição geográfica podem mudar.",
        "A comparação de TTLs representa estados de cache diferentes, não necessariamente políticas fixas dos provedores.",
        "Consultas diretas a resolvedores públicos dependem de UDP/TCP porta 53 e podem sofrer bloqueio por VPN ou rede institucional.",
        "Os resultados reproduzem uma execução controlada; repetições devem registrar data e valores observados.",
    ])

    add_heading("Conclusão", page_before=True)
    add_body(
        "O experimento confirmou a estrutura de mensagens DNS, a correlação por Transaction ID, o papel das flags e do TTL, e as diferenças entre tipos de registro. "
        "Também demonstrou que consultas equivalentes podem apresentar TTLs distintos por causa do cache, que o servidor selecionado altera o destino real dos pacotes e que EDNS0 permite respostas UDP superiores a 512 bytes. "
        "O ambiente Podman tornou o caminho cliente–resolvedor observável e permitiu limpar o cache do Unbound, repetir consultas e validar UDP e TCP de forma controlada."
    )

    add_band("Referências")
    add_list([
        "monitor_dns_v2.docx — roteiro do experimento executado.",
        "RFC 1034 e RFC 1035 — conceitos, formato e operação do DNS.",
        "RFC 6891 — Extension Mechanisms for DNS (EDNS0).",
        "Wireshark User's Guide; documentação do Unbound e do Podman.",
    ])

    section = doc.sections[0]
    section.left_margin = Inches(1.18)
    section.right_margin = Inches(1.18)
    section.top_margin = Inches(0.98)
    section.bottom_margin = Inches(0.98)

    for drawing_property in doc._element.iter(qn("wp:docPr")):
        drawing_property.set("title", "Brasão da PUCRS")
        drawing_property.set(
            "descr", "Brasão da Pontifícia Universidade Católica do Rio Grande do Sul"
        )

    doc.core_properties.title = "Relatório de solução — Monitoração do DNS"
    doc.core_properties.subject = "Resultados dos experimentos DNS no ambiente Podman"
    doc.core_properties.comments = "Solução baseada em execução realizada em 24/08/2026."
    WORKING.parent.mkdir(parents=True, exist_ok=True)
    doc.save(WORKING)
    compose_fidelity_package(SOURCE, WORKING, OUTPUT)

    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("O documento de referência foi alterado durante a geração.")
    print(OUTPUT)


if __name__ == "__main__":
    build()
